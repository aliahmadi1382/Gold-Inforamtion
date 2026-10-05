"""Read-only Pink Sheet monthly gold adapter; never a daily-price substitute."""

import io
import math
import re
import zipfile
from datetime import UTC, date, datetime
from itertools import islice

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .ingestion import fetch_bytes, provenance
from .models import Observation, aware

MONTHLY_URL = (
    "https://thedocs.worldbank.org/en/doc/"
    "74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx"
)
GOLD_DESCRIPTION = (
    "Gold, spot average of daily rates, from June 2025; previously (UK), 99.5% fine, "
    "London afternoon fixing, average of daily rates"
)
DATASET = "pink_sheet_gold_monthly_v1"
SERIES = "WB_GOLD_MONTHLY"


def parse_monthly_gold(content, source, digest, retrieved) -> list[Observation]:
    retrieved = aware(retrieved).astimezone(UTC)
    if source.source_id != "world_bank_pink_sheet":
        raise ValueError("monthly adapter requires the World Bank registry source")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if (
                len(archive.infolist()) > 500
                or sum(f.file_size for f in archive.infolist()) > 30_000_000
            ):
                raise ValueError("workbook exceeds expanded size limit")
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError("invalid Pink Sheet workbook") from exc
    try:
        if not {"Monthly Prices", "Description"} <= set(workbook.sheetnames):
            raise ValueError("missing Pink Sheet sheets")
        sheet = workbook["Monthly Prices"]
        descriptions = workbook["Description"]
        if (
            sheet.max_row > 5000
            or sheet.max_column > 200
            or descriptions.max_row > 1000
            or descriptions.max_column > 200
        ):
            raise ValueError("unexpected Pink Sheet dimensions")
        gold_descriptions = [
            str(row[1]).strip()
            for row in descriptions.iter_rows(values_only=True)
            if len(row) > 1 and isinstance(row[1], str) and row[1].strip().startswith("Gold,")
        ]
        if gold_descriptions != [GOLD_DESCRIPTION]:
            raise ValueError("gold methodology changed; review before importing")
        rows = sheet.iter_rows(values_only=True)
        heading = list(islice(rows, 6))
        if len(heading) != 6:
            raise ValueError("missing monthly headings")
        columns = [i for i, value in enumerate(heading[4]) if value == "Gold"]
        if len(columns) != 1 or heading[5][columns[0]] != "($/troy oz)":
            raise ValueError("unexpected gold column or unit")
        column = columns[0]
        records = []
        previous = None
        for row_number, row in enumerate(rows, 7):
            label = row[0]
            if not isinstance(label, str) or not re.fullmatch(r"\d{4}M\d{2}", label):
                raise ValueError("invalid monthly period label")
            period = date(int(label[:4]), int(label[-2:]), 1)
            serial = period.year * 12 + period.month
            if previous is not None and serial != previous + 1:
                raise ValueError("duplicate, unordered or missing monthly period")
            previous = serial
            if period >= retrieved.date().replace(day=1):
                raise ValueError("monthly average includes an unfinished or future month")
            value = row[column]
            if (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not math.isfinite(value)
                or value <= 0
            ):
                raise ValueError("invalid or missing monthly gold value")
            p = provenance(
                source,
                digest,
                f"Monthly Prices!{get_column_letter(column + 1)}{row_number}",
                datetime.combine(period, datetime.min.time(), UTC),
                retrieved,
                DATASET,
                "currency_per_troy_ounce",
                "USD",
                label,
                transformations=("pink_sheet_xlsx_v1", "month_start_period_label_utc"),
            )
            p = type(p).model_validate({**p.model_dump(), "source_url": MONTHLY_URL})
            records.append(
                Observation(
                    provenance=p,
                    layer="price",
                    series_id=SERIES,
                    value=float(value),
                    dimensions={
                        "instrument": "GOLD",
                        "frequency": "1mo",
                        "aggregation": "monthly_average",
                        "methodology": "spot_daily_average"
                        if period >= date(2025, 6, 1)
                        else "london_afternoon_fixing_average",
                    },
                )
            )
        if not records:
            raise ValueError("empty monthly gold history")
        return records
    finally:
        workbook.close()


def ingest_monthly_gold(store, source, fetch=fetch_bytes) -> int:
    content = fetch(MONTHLY_URL)
    retrieved = datetime.now(UTC)
    return store.put(parse_monthly_gold(content, source, store.put_raw(content), retrieved))
