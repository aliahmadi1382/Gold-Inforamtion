"""Native annualized CPI release estimates; never measured release-time evidence."""

import io
import math
import re
import zipfile
from datetime import UTC, datetime

from openpyxl import load_workbook

from .ingestion import fetch_bytes, provenance
from .models import Observation, aware

SOURCE = "philadelphia_fed_pcpi"
DATASET = "philly_pcpi_release_estimates_v1"
URL = (
    "https://www.philadelphiafed.org/-/media/FRBP/Assets/Surveys-And-Data/"
    "real-time-data/data-files/xlsx/pcpi_first_second_third.xlsx"
    "?sc_lang=en&hash=3AD2A1024E65EFCDA2FF94C017C540E6"
)
UNIT = "percent_change_mom_annualized_sa"
STAGES = ("First", "Second", "Third", "Most_Recent")


def parse_pcpi(content, source, digest, retrieved):
    if source.source_id != SOURCE:
        raise ValueError("PCPI adapter requires the Philadelphia Fed source")
    retrieved = aware(retrieved).astimezone(UTC)
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if (
                len(archive.infolist()) > 500
                or sum(entry.file_size for entry in archive.infolist()) > 30_000_000
            ):
                raise ValueError("workbook exceeds expanded size limit")
    except zipfile.BadZipFile as exc:
        raise ValueError("invalid PCPI workbook") from exc
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    try:
        if set(workbook.sheetnames) != {"NOTES", "DATA"}:
            raise ValueError("unexpected PCPI sheets")
        notes, sheet = workbook["NOTES"], workbook["DATA"]
        if (
            notes.max_row > 100
            or notes.max_column != 1
            or sheet.max_row > 5000
            or sheet.max_row < 6
            or sheet.max_column != 5
        ):
            raise ValueError("unexpected PCPI dimensions")
        descriptions = [r[0] for r in notes.iter_rows(values_only=True)]
        required = {
            "Variable ID: PCPI",
            "Unit of Measurement: M/M Growth (Annual Rate, Percentage Points)",
            "Source: Bureau of Labor Statistics (BLS)",
        }
        if not required <= set(descriptions):
            raise ValueError("PCPI metric/unit/source changed; review required")
        rows = sheet.iter_rows(values_only=True)
        heading = [next(rows, None) for _ in range(5)]
        if (
            heading[4] != ("Date", *STAGES)
            or heading[1][0] != "M/M Growth (Annual Rate, Percentage Points)"
        ):
            raise ValueError("unexpected PCPI header")
        records, previous = [], None
        for row_number, row in enumerate(rows, 6):
            label = row[0]
            if not isinstance(label, str) or not re.fullmatch(r"\d{4}:(?:0[1-9]|1[0-2])", label):
                raise ValueError("invalid PCPI reference month")
            observed = datetime.strptime(label, "%Y:%m").replace(tzinfo=UTC)
            if (previous is not None and observed <= previous) or observed > retrieved:
                raise ValueError("duplicate, unordered or future PCPI month")
            previous = observed
            for column, stage in enumerate(STAGES, 1):
                value = row[column]
                missing = value is None or value == "#N/A"
                if not missing and (
                    isinstance(value, bool)
                    or not isinstance(value, int | float)
                    or not math.isfinite(value)
                    or value <= -100
                ):
                    raise ValueError("invalid PCPI growth value")
                p = provenance(
                    source,
                    digest,
                    f"DATA:{row_number}:{column + 1}",
                    observed,
                    retrieved,
                    DATASET,
                    UNIT,
                    None,
                    str(value),
                    transformations=("native annualized M/M growth retained",),
                )
                records.append(
                    Observation(
                        provenance=p,
                        layer="macro",
                        series_id="PHILLY_PCPI_" + stage.upper(),
                        value=None if missing else float(value),
                        dimensions=dict(
                            frequency="M",
                            release_stage=stage,
                            seasonal_adjustment="SA",
                            metric="cpi_all_items",
                            estimation_basis="monthly_vintage_algorithm",
                            release_time_evidence="not_measured",
                        ),
                    )
                )
        if not records:
            raise ValueError("empty PCPI history")
        return records
    finally:
        workbook.close()


def ingest_pcpi(store, source, fetch=fetch_bytes):
    content = fetch(URL)
    retrieved = datetime.now(UTC)
    return store.put(parse_pcpi(content, source, store.put_raw(content), retrieved))
